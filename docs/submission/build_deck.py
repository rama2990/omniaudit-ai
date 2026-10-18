"""Fill the official AI Builder Cup 2026 template with OmniAudit content.

    cd backend && uv run --with python-pptx python ../docs/submission/build_deck.py [video_url]
Output: docs/submission/OmniAudit_AI_Builder_Cup_2026.pptx (convert to PDF with PowerPoint).
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
VIDEO = sys.argv[1] if len(sys.argv) > 1 else "[demo video link]"
APP = "https://omniaudit-648878449957.asia-southeast2.run.app"
REPO = "https://github.com/rama2990/omniaudit-ai"
FONT, FONT_M = "Google Sans Flex", "Google Sans Flex Medium"
INK, MUTED, BLUE = RGBColor(0x20, 0x21, 0x24), RGBColor(0x3C, 0x40, 0x43), RGBColor(0x1A, 0x73, 0xE8)
GREEN, RED, AMBER = RGBColor(0x18, 0x80, 0x38), RGBColor(0xC5, 0x22, 0x1F), RGBColor(0xB0, 0x60, 0x00)
SOFT = {"blue": RGBColor(0xE8, 0xF0, 0xFE), "green": RGBColor(0xE6, 0xF4, 0xEA), "amber": RGBColor(0xFE, 0xF7, 0xE0),
        "red": RGBColor(0xFC, 0xE8, 0xE6), "grey": RGBColor(0xF1, 0xF3, 0xF4)}

prs = Presentation(HERE / "template.pptx")
S = list(prs.slides)


def text(slide, x, y, w, h, paras, size=12, anchor=MSO_ANCHOR.TOP):
    """paras: list of str or (str, dict(bold, color, size, font, bullet))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = o.get("align", PP_ALIGN.LEFT)
        para.space_after = Pt(o.get("after", 4))
        r = para.add_run()
        r.text = ("•  " if o.get("bullet") else "") + t
        f = r.font
        f.size = Pt(o.get("size", size))
        f.name = o.get("font", FONT_M if o.get("bold") else FONT)
        f.bold = False
        f.color.rgb = o.get("color", MUTED if not o.get("bold") else INK)
    return tb


def box(slide, x, y, w, h, title, body="", fill="blue", tsize=11, bsize=9, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = SOFT[fill]
    sh.line.color.rgb = RGBColor(0xDA, 0xDC, 0xE0)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.06))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size, r.font.name, r.font.color.rgb = Pt(tsize), FONT_M, INK
    if body:
        q = tf.add_paragraph()
        q.alignment = PP_ALIGN.CENTER
        r = q.add_run()
        r.text = body
        r.font.size, r.font.name, r.font.color.rgb = Pt(bsize), FONT, MUTED
    return sh


def arrow(slide, x, y, w=0.28, h=0.22, down=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW if down else MSO_SHAPE.RIGHT_ARROW,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x9A, 0xA0, 0xA6)
    sh.line.fill.background()
    return sh


def image(slide, name, x, y, w=None, h=None):
    f = HERE / "img" / name
    if f.exists():
        pic = slide.shapes.add_picture(str(f), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
        pic.line.color.rgb = RGBColor(0xDA, 0xDC, 0xE0)
        return pic
    text(slide, x, y, w or 4, 0.4, [(f"[screenshot {name}]", {"color": RED})])


# 1. Team details
for para in S[0].shapes[0].text_frame.paragraphs:
    t = para.text.strip()
    val = {"Team name:": "Team OmniAudit", "Team leader name:": "Rama Raditya",
           "Problem Statement:": "Future of Work & Enterprise Productivity: finance teams at multi-entity "
                                 "property developers audit contractor bills and match buyer payments by hand"}.get(t)
    if val:
        r = para.add_run()
        r.text = " " + val
        r.font.size, r.font.name = Pt(16 if t != "Problem Statement:" else 13), FONT

# 3. Brief about the idea
text(S[2], 0.4, 1.35, 9.2, 1.3, [
    ("OmniAudit AI is an AI audit co-pilot for property developers that run every project as a separate company (SPV).", {"bold": True, "size": 14}),
    ("It checks each contractor bill before money goes out and matches each buyer payment to its unit when money comes in. "
     "Gemini reads the documents, deterministic code calculates every rupiah, and people make the decisions. Clean items are "
     "approved in one click; suspicious ones are locked and opened in a split-view audit screen, and every action lands in a "
     "tamper-evident audit trail.", {"size": 12}),
])
for i, (t, b, c) in enumerate([("Gemini reads", "Pages classified and fields extracted; every value must point to a word in the document", "blue"),
                               ("Code calculates", "VAT, construction withholding tax, retention, down payment, late penalties, journals", "green"),
                               ("People decide", "One-click approval for clean bills; disputes, CFO approvals and waivers with reasons", "amber")]):
    box(S[2], 0.4 + i * 3.1, 3.35, 2.9, 1.6, t, b, c, tsize=14, bsize=10)
    if i < 2:
        arrow(S[2], 3.32 + i * 3.1, 4.05, 0.16, 0.2)

# 4. Opportunities
tf = S[3].shapes[3].text_frame
tf.clear()
qa = [("Opportunities", "Large developers process thousands of contractor bills and tens of thousands of buyer transfers a month. "
                        "Overbilling, missed retention or penalties, and unallocated receipts leak real money and take days to check by hand."),
      ("How different is it from other ideas?", "Generic invoice OCR stops at extraction. OmniAudit grounds every AI value in the source document "
                                                "and runs Indonesian construction rules (PP 9/2022 withholding by contractor grade, Coretax VAT base, "
                                                "retention, down payment, penalties) in deterministic code, across many SPVs at once."),
      ("How will it solve the problem?", "Eight automatic checks lock risky bills before payment; a match score assigns transfers to units; "
                                         "the CFO gets one queue for large or suspicious cases; every decision is hash-chained."),
      ("USP", "AI never calculates money. Results are explainable to an external auditor: each number links to its location in the PDF.")]
for i, (q, a) in enumerate(qa):
    for j, (t, size, font, col) in enumerate([(q, 13, FONT_M, INK), (a, 11, FONT, MUTED)]):
        para = tf.paragraphs[0] if i == 0 and j == 0 else tf.add_paragraph()
        para.space_after = Pt(2 if j == 0 else 8)
        r = para.add_run()
        r.text = t
        r.font.size, r.font.name, r.font.color.rgb = Pt(size), font, col
S[3].shapes[3].height = Inches(4.6)

# 5. Features
feat_ap = ["Upload a multi-page PDF: invoice, inspection report (BAP), tax invoice, photos",
           "Local OCR (digital and scanned) + Gemini extraction with grounding",
           "Contract agent: links the bill to its contract, vendor grade and terms",
           "8 anomaly checks: overbilling, duplicate across SPVs, inflated amount, arithmetic, missing retention, wrong tax rate, tax-invoice mismatch, invalid progress",
           "Two-column calculation: claimed vs inspection-verified, to the rupiah",
           "Split-view audit screen with highlight boxes and click-to-jump evidence",
           "Overrides with mandatory reasons; balanced journal posted to ERP"]
feat_ar = ["Bank statement CSV upload; exact virtual-account matches auto-allocated",
           "Match score: 45% unit code + 35% payer name + 20% amount",
           "Late fee 0.1%/day per installment; waivers with reasons, executive waivers to CFO",
           "FIFO allocation, buyer credit, bank-fee tolerance",
           "Receipt PDF + WhatsApp/email notification (simulated)"]
feat_gov = ["Role-based access and data isolation per SPV", "CFO room: approvals above IDR 500 million, hard stops",
            "Hash-chained audit trail with integrity check", "AI accuracy tracked per field from human corrections",
            "Dashboard: unallocated payments, pending bills, disputes, leakage prevented"]
for i, (title, items) in enumerate([("Payables (contractor bills)", feat_ap), ("Receivables (buyer payments)", feat_ar),
                                    ("Control & governance", feat_gov)]):
    text(S[4], 0.4 + i * 3.1, 1.35, 2.95, 4.1, [(title, {"bold": True, "size": 13, "color": BLUE})] +
         [(t, {"bullet": True, "size": 9.5, "after": 3}) for t in items])

# 6. Process flow
y1, y2 = 1.55, 3.75
text(S[5], 0.4, 1.3, 4, 0.3, [("Contractor bill (payables)", {"bold": True, "size": 11, "color": BLUE})])
ap = [("Upload PDF", "AP lead", "grey"), ("Local OCR", "words + coordinates", "grey"), ("Gemini on Vertex", "classify & extract, grounded", "blue"),
      ("Contract agent", "terms, vendor grade", "blue"), ("Auditor rules", "8 checks", "red"), ("Calculator", "tax, retention, penalty", "green")]
for i, (t, b, c) in enumerate(ap):
    box(S[5], 0.4 + i * 1.56, y1 + 0.25, 1.3, 0.85, t, b, c, tsize=10, bsize=8)
    if i < len(ap) - 1:
        arrow(S[5], 1.72 + i * 1.56, y1 + 0.57, 0.2, 0.2)
outs = [("Ready", "one-click approve", "green"), ("Dispute", "split-view review", "amber"), ("Hard stop", "CFO only", "red")]
for i, (t, b, c) in enumerate(outs):
    box(S[5], 3.2 + i * 1.5, y1 + 1.3, 1.35, 0.6, t, b, c, tsize=10, bsize=8)
arrow(S[5], 8.25, y1 + 1.15, 0.2, 0.25, down=True)
box(S[5], 7.95, y1 + 1.4, 1.6, 0.5, "Post to ERP", "balanced journal", "blue", tsize=10, bsize=8)
text(S[5], 0.4, y2 - 0.1, 4, 0.3, [("Buyer payment (receivables)", {"bold": True, "size": 11, "color": BLUE})])
arf = [("Bank CSV", "or VA feed", "grey"), ("Match score", "unit, name, amount", "blue"), ("Late fee", "0.1%/day, waiver", "amber"),
       ("Allocate", "FIFO, buyer credit", "green"), ("Journal + receipt", "WhatsApp / email", "blue")]
for i, (t, b, c) in enumerate(arf):
    box(S[5], 0.4 + i * 1.9, y2 + 0.25, 1.55, 0.8, t, b, c, tsize=10, bsize=8)
    if i < len(arf) - 1:
        arrow(S[5], 2.0 + i * 1.9, y2 + 0.55, 0.22, 0.2)
text(S[5], 0.4, 4.95, 9.2, 0.35, [("Every step writes to the hash-chained audit trail (mirrored to Firestore).", {"size": 10})])

# 7. Wireframes -> actual screens
image(S[6], "ap-review.png", 0.4, 1.35, h=4.0)
text(S[6], 7.6, 1.4, 2.0, 3.9, [("Split-view audit", {"bold": True, "size": 12}),
                                ("Left: original PDF with highlight boxes from OCR coordinates; red marks values behind a finding.", {"size": 10}),
                                ("Right: claimed vs verified amounts, findings, and the live agent trail.", {"size": 10})])

# 8. Architecture
box(S[7], 0.4, 2.55, 1.6, 1.0, "Browser", "React + Material UI\npersona & SPV switch", "grey", 11, 8)
arrow(S[7], 2.08, 2.95, 0.3, 0.22)
run = S[7].shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(1.4), Inches(3.9), Inches(3.75))
run.fill.solid(); run.fill.fore_color.rgb = RGBColor(0xFD, 0xFD, 0xFC); run.line.color.rgb = BLUE
text(S[7], 2.6, 1.45, 3.7, 0.3, [("Google Cloud Run (asia-southeast2)", {"bold": True, "size": 11, "color": BLUE})])
box(S[7], 2.65, 1.85, 3.6, 0.6, "FastAPI service + React static build", "one container, always-on CPU", "blue", 10, 8)
box(S[7], 2.65, 2.55, 1.75, 0.75, "Agent pipeline", "parser, contract,\nauditor, calculator", "blue", 10, 8)
box(S[7], 4.5, 2.55, 1.75, 0.75, "Local OCR", "PyMuPDF + RapidOCR\n(PaddleOCR models)", "grey", 10, 8)
box(S[7], 2.65, 3.4, 1.75, 0.75, "Rule & tax engine", "Decimal, deterministic", "green", 10, 8)
box(S[7], 4.5, 3.4, 1.75, 0.75, "Demo data store", "SQLite per instance,\nreset snapshots", "grey", 10, 8)
box(S[7], 2.65, 4.25, 3.6, 0.7, "Mock integrations", "ERP posting, WhatsApp/email receipts, bank feed", "grey", 10, 8)
for i, (t, b, c) in enumerate([("Vertex AI", "Gemini 2.5 Flash\nstructured JSON output", "blue"),
                               ("Firestore", "audit trail & ERP\njournal mirror", "amber"),
                               ("Cloud Build + Artifact Registry", "container build & deploy", "grey")]):
    arrow(S[7], 6.5, 1.85 + i * 1.15 + 0.3, 0.3, 0.22)
    box(S[7], 6.9, 1.6 + i * 1.15, 2.7, 0.9, t, b, c, 11, 8)
text(S[7], 0.4, 3.75, 1.9, 1.2, [("Service account auth to Vertex AI and Firestore; no API keys on the server.", {"size": 9})])

# 9. Technologies
rows = [("Gemini 2.5 Flash on Vertex AI", "Page classification and field extraction with a JSON schema; values grounded to OCR words"),
        ("Google Cloud Run", "Single container for API + web app, Jakarta region, always-on CPU for background processing"),
        ("Cloud Firestore", "Durable mirror of the audit trail and posted ERP journals"),
        ("Cloud Build / Artifact Registry", "Source-to-container builds for each deploy"),
        ("PyMuPDF + RapidOCR", "Local OCR with word coordinates for evidence boxes; page images never leave the server"),
        ("FastAPI, SQLModel, Python Decimal", "Agent pipeline, rule engine and tax calculator, 77 automated tests"),
        ("React, Vite, Material UI, pdf.js", "Dashboard, split-view audit screen, reconciliation and CFO room")]
tbl = S[8].shapes.add_table(len(rows), 2, Inches(0.4), Inches(1.4), Inches(9.2), Inches(3.8)).table
tbl.columns[0].width, tbl.columns[1].width = Inches(2.9), Inches(6.3)
for r, (a, b) in enumerate(rows):
    for c, val in enumerate((a, b)):
        cell = tbl.cell(r, c)
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size, para.runs[0].font.name = Pt(10.5), FONT_M if c == 0 else FONT
        para.runs[0].font.color.rgb = INK if c == 0 else MUTED
        cell.fill.solid()
        cell.fill.fore_color.rgb = SOFT["blue"] if c == 0 else RGBColor(0xFD, 0xFD, 0xFC)

# 10. Cost (estimate)
text(S[9], 0.4, 1.35, 9.2, 3.9, [
    ("Prototype run cost (estimates, Google Cloud list prices)", {"bold": True, "size": 13}),
    ("Cloud Run, 1 instance, 2 vCPU / 2 GiB, always-on CPU: about IDR 1.7 million per month", {"bullet": True}),
    ("Gemini 2.5 Flash on Vertex AI: one request per bill (a few thousand tokens); cost scales with bill volume and is small per document", {"bullet": True}),
    ("Firestore: a few writes per action, well inside low-volume pricing", {"bullet": True}),
    ("Local OCR runs inside the container: no per-page OCR fees", {"bullet": True}),
    ("Production", {"bold": True, "size": 13}),
    ("Autoscaling Cloud Run, Vertex AI regional endpoint for data residency, and managed database replacing the demo SQLite store; "
     "cost driven mainly by document volume, not by user count.", {"size": 11}),
], size=11)

# 11. Snapshots
image(S[10], "dashboard.png", 0.4, 1.35, h=1.95)
image(S[10], "ar-review.png", 3.85, 1.35, h=1.95)
image(S[10], "cfo.png", 0.4, 3.4, h=1.95)
text(S[10], 7.3, 1.4, 2.3, 3.9, [("Dashboard", {"bold": True, "size": 11}), ("Unallocated payments, pending bills, disputes, leakage prevented.", {"size": 9}),
                                 ("Reconciliation", {"bold": True, "size": 11}), ("97% match to unit C1-02, late fee waived for a pending mortgage.", {"size": 9}),
                                 ("CFO room", {"bold": True, "size": 11}), ("Approvals, hard stops, audit trail integrity, AI accuracy.", {"size": 9})])

# 12. Performance
stats = [("15 / 15", "full-run scenario steps pass on the live Cloud Run site with Gemini 2.5 Flash (12 bills + 3 bank statements)"),
         ("98 / 98", "fields matched in the Gemini extraction check on the demo documents"),
         ("15–25 s", "per digital bill end to end on Cloud Run; about 50 s for a scanned bill with local OCR"),
         ("77", "automated tests pass, including tax and journal examples that match to the rupiah"),
         ("IDR 540.8M", "overpayment prevented in the simulated work week (claimed minus verified)")]
for i, (n, d) in enumerate(stats):
    col, row = i % 2, i // 2
    box(S[11], 0.4 + col * 4.65, 1.4 + row * 1.3, 4.45, 1.15, n, d, ["green", "blue", "blue", "blue", "green"][i], 18, 9.5)

# 13. Future
text(S[12], 0.4, 1.35, 4.5, 3.9, [("Next steps to production", {"bold": True, "size": 13})] + [
    (t, {"bullet": True, "size": 11}) for t in ["Real ERP connectors (SAP, Oracle)", "MT940 bank statements and bank webhooks",
                                                 "WhatsApp Business API for receipts", "Vertex AI regional endpoint for data residency",
                                                 "Retention release at final handover (BAST II)", "Pilot in one SPV with real data"]])
text(S[12], 5.1, 1.35, 4.5, 3.9, [("Impact targets", {"bold": True, "size": 13})] + [
    (t, {"bullet": True, "size": 11}) for t in ["Audit a contractor bill in under 15 minutes instead of 3–5 working days",
                                                 "No overbilling slips through; retention, penalties and down payment always deducted",
                                                 "95% of buyer payments allocated the same day"]])

# 14. Links
for para in S[13].shapes[3].text_frame.paragraphs:
    t = para.text.strip()
    val = {"GitHub Public Repository": REPO, "Demo Video Link (3 Minutes)": VIDEO, "Final Product Link": APP}.get(t)
    if val:
        r = para.add_run()
        r.text = ": " + val
        r.font.size, r.font.name, r.font.color.rgb = Pt(12), FONT, BLUE
        r.hyperlink.address = val if val.startswith("http") else None

# 15. Worked example
text(S[14], 0.3, 0.7, 9.6, 0.5, [("One bill from the demo", {"bold": True, "size": 16, "color": INK})])
text(S[14], 0.4, 1.35, 9.2, 0.6, [("Magnolia Cluster contract, progress payment 3: the invoice claims 80% progress, the site inspection says 75%, "
                                   "and the project is 10 days late.", {"size": 12})])
box(S[14], 0.4, 2.15, 2.9, 1.2, "IDR 2,575,450,000", "requested by the contractor", "grey", 17, 10)
box(S[14], 3.55, 2.15, 2.9, 1.2, "IDR 2,112,875,000", "after verification (progress, down payment, retention, late penalty, withholding tax)", "blue", 17, 9)
box(S[14], 6.7, 2.15, 2.9, 1.2, "IDR 462,575,000", "overpayment stopped before payment", "green", 17, 10)
text(S[14], 0.4, 3.6, 9.2, 1.6, [("Because the claim is above IDR 500 million, the bill is locked and routed to the CFO. "
                                  "The auditor sees the 80% and 75% highlighted in the original PDF, accepts the correction, and the CFO approves; "
                                  "a balanced journal is posted to ERP and the decision is logged in the hash-chained audit trail.", {"size": 11})])

# remove the template's instruction slide (slide 2)
sld = prs.slides._sldIdLst
sld.remove(list(sld)[1])

out = HERE / "OmniAudit_AI_Builder_Cup_2026.pptx"
prs.save(out)
print(out)
